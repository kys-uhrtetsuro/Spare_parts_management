import io
import json
from django.test import TestCase

from .models import EquipmentMaster, EquipmentSection, PartCategory, PartMaster, PartVersion, StockAlertSetting, StockHistory, SystemMaster
from .services import InputError, change_part, correct_movement, create_movement, inventory_rows, replace_alerts


class InventoryTests(TestCase):
    def setUp(self):
        self.system = SystemMaster.objects.create(name="制御A")
        self.other = SystemMaster.objects.create(name="制御B")
        self.part = PartMaster.objects.create(part_number="P100", model_number="V1", name="センサ")
        self.actor = {"operator": "担当者"}

    def movement(self, direction, quantity, system=None, part=None):
        return create_movement({
            "system_id": (system or self.system).id,
            "part_id": (part or self.part).id, "direction": direction,
            "quantity": quantity, **self.actor,
        })

    def test_company_is_not_required_for_movement_or_correction(self):
        first = self.movement("IN", 2)
        self.assertFalse(hasattr(StockHistory.objects.get(pk=first.pk), "company"))
        corrected = correct_movement(first.pk, {
            "direction": "IN", "quantity": 1,
            "note": "記載修正", "operator": "担当者",
        })
        self.assertEqual(corrected[1].operator, "担当者")
        self.assertEqual(inventory_rows()[0]["quantity"], 1)

    def test_system_separation_and_negative_block(self):
        self.movement("IN", 5)
        self.movement("OUT", 2)
        self.movement("IN", 4, system=self.other)
        rows = inventory_rows()
        self.assertEqual({r["system"]: r["quantity"] for r in rows}, {"制御A": 3, "制御B": 4})
        with self.assertRaises(InputError):
            self.movement("OUT", 4)
        self.assertEqual(StockHistory.objects.count(), 3)

    def test_zero_is_only_exported_if_monitored(self):
        self.movement("IN", 1)
        self.movement("OUT", 1)
        self.assertFalse(inventory_rows())
        replace_alerts(self.system, [{"part_id": self.part.id, "threshold": 2}])
        self.assertEqual(inventory_rows()[0]["quantity"], 0)
        self.assertTrue(inventory_rows()[0]["alert"])
        replace_alerts(self.system, [])
        self.assertFalse(inventory_rows())

    def test_positive_alerts_atomic(self):
        replace_alerts(self.system, [{"part_id": self.part.id, "threshold": 3}])
        with self.assertRaises(InputError):
            replace_alerts(self.system, [{"part_id": self.part.id, "threshold": 0}])
        self.assertEqual(StockAlertSetting.objects.get().threshold, 3)

    def test_successor_keeps_old_identifier_and_its_stock(self):
        old = self.movement("IN", 3)
        successor = change_part(self.part, {
            "part_number": "P200", "model_number": "V2", "name": "センサ2",
        })
        self.assertNotEqual(successor.id, self.part.id)
        self.assertEqual(successor.previous_part_id, self.part.id)
        self.assertEqual((self.part.part_number, self.part.model_number), ("P100", "V1"))
        history = StockHistory.objects.get(pk=old.pk)
        self.assertEqual((history.part_id, history.part_number_snapshot), (self.part.id, "P100"))
        self.assertEqual(inventory_rows()[0]["part_number"], "P100")
        self.movement("IN", 2, part=successor)
        self.assertEqual(
            {r["part_number"]: r["quantity"] for r in inventory_rows()},
            {"P100": 3, "P200": 2},
        )
        self.assertEqual(PartVersion.objects.get().part_id, self.part.id)
        with self.assertRaises(InputError):
            self.movement("OUT", 4, part=successor)

    def test_model_only_and_double_change_linkage(self):
        original = self.movement("IN", 8)
        intermediate = change_part(self.part, {
            "part_number": "P100", "model_number": "V2", "name": "センサ",
        })
        final = change_part(intermediate, {
            "part_number": "P200", "model_number": "V3", "name": "センサ改",
        })
        self.assertEqual(intermediate.previous_part_id, self.part.pk)
        self.assertEqual(final.previous_part_id, intermediate.pk)
        self.assertEqual(StockHistory.objects.get(pk=original.pk).model_number_snapshot, "V1")
        self.assertEqual(inventory_rows()[0]["model_number"], "V1")
        self.assertEqual(inventory_rows()[0]["quantity"], 8)
        self.assertEqual(PartVersion.objects.count(), 2)
        response = self.client.get(f"/api/parts/{final.pk}/versions")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["items"]), 2)

    def test_machine_section_category_reverse_lookup(self):
        machine = EquipmentMaster.objects.create(name="搬送装置")
        section = EquipmentSection.objects.create(equipment=machine, name="制御盤")
        category = PartCategory.objects.create(name="制御機器")
        self.part.equipment = machine
        self.part.section = section
        self.part.category = category
        self.part.save()
        another = PartMaster.objects.create(
            part_number="P500", model_number="V1", name="ベアリング",
        )
        self.movement("IN", 2)
        self.movement("IN", 3, part=another)
        params = f"?equipment={machine.pk}&section={section.pk}&category={category.pk}"
        result = self.client.get("/api/parts" + params)
        self.assertEqual(result.status_code, 200)
        self.assertEqual([row["id"] for row in result.json()["items"]], [self.part.pk])
        rows = self.client.get("/api/stocks" + params).json()["items"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["part_id"], self.part.pk)
        from openpyxl import load_workbook
        excel = self.client.get("/api/export" + params)
        ws = load_workbook(io.BytesIO(b"".join(excel.streaming_content))).active
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["B2"].value, "P100")

    def test_correction_is_append_only_and_cannot_breach_zero(self):
        original = self.movement("IN", 5)
        self.movement("OUT", 3)
        with self.assertRaises(InputError):
            correct_movement(original.pk, {"direction": "IN", "quantity": 2, "note": "誤記", **self.actor})
        self.assertEqual(StockHistory.objects.count(), 2)
        reverse, revised = correct_movement(
            original.pk, {"direction": "IN", "quantity": 4, "note": "誤記", **self.actor}
        )
        self.assertEqual(reverse.reversal_of_id, original.pk)
        self.assertEqual(revised.quantity, 4)
        self.assertEqual(inventory_rows()[0]["quantity"], 1)
        with self.assertRaises(InputError):
            correct_movement(original.pk, {"direction": "IN", "quantity": 4, "note": "重複", **self.actor})

    def test_excel_roundtrip_forks_system_stock_and_metadata(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        from .models import StockImportBatch
        machine = EquipmentMaster.objects.create(name="仕分け装置", alias="ソータ")
        section = EquipmentSection.objects.create(equipment=machine, name="搬送系")
        category = PartCategory.objects.create(name="制御機器")
        self.part.equipment, self.part.section, self.part.category = machine, section, category
        self.part.save()
        version = change_part(self.part, {
            "part_number": "P200", "model_number": "V2", "name": "センサ新",
        })
        self.movement("IN", 8)
        self.movement("IN", 3, part=version)
        self.movement("OUT", 1, part=version)
        replace_alerts(self.system, [
            {"part_id": version.id, "threshold": 3}
        ])

        download = self.client.get(f"/api/export?system={self.system.pk}")
        self.assertEqual(download.status_code, 200)
        content = b"".join(download.streaming_content)
        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book["__systems"].sheet_state, "hidden")
        self.assertEqual(book["__parts"].sheet_state, "hidden")
        self.assertEqual(book["__equipment"].sheet_state, "hidden")
        self.assertEqual(book["__sections"].sheet_state, "hidden")
        self.assertEqual(book["__categories"].sheet_state, "hidden")
        self.assertEqual(book["現在庫"]["E2"].value, 8)
        book.close()

        def upload(path, **fields):
            return self.client.post(
                path, data={"file": SimpleUploadedFile(
                    "移植元.xlsx", content,
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ), **fields}
            )

        preview = upload("/api/import/preview")
        self.assertEqual(preview.status_code, 200, preview.content)
        self.assertEqual(preview.json()["part_count"], 2)
        self.assertEqual(preview.json()["system_count"], 1)
        self.assertEqual(preview.json()["quantity_total"], 10)
        self.assertIn("制御A", preview.json()["conflicts"])

        duplicate_name = upload("/api/import/apply")
        self.assertEqual(duplicate_name.status_code, 400)
        self.assertEqual(StockImportBatch.objects.count(), 0)

        applied = upload("/api/import/apply", system_name="制御A-複製")
        self.assertEqual(applied.status_code, 201, applied.content)
        new = SystemMaster.objects.get(name="制御A-複製")
        from .services import balance
        self.assertEqual(balance(new.pk, self.part.pk), 8)
        self.assertEqual(balance(new.pk, version.pk), 2)
        self.assertEqual(
            StockAlertSetting.objects.get(system=new, part=version).threshold, 3
        )
        self.assertEqual(PartVersion.objects.count(), 1)
        self.assertEqual(StockHistory.objects.filter(system=new).count(), 2)

        twice = upload("/api/import/apply", system_name="制御A-複製2")
        self.assertEqual(twice.status_code, 400)
        self.assertFalse(SystemMaster.objects.filter(name="制御A-複製2").exists())
        self.assertEqual(StockImportBatch.objects.count(), 1)

    def test_excel_snapshot_allows_quantity_and_threshold_edits(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        self.movement("IN", 2)
        book = load_workbook(io.BytesIO(
            b"".join(self.client.get(f"/api/export?system={self.system.pk}").streaming_content)
        ))
        book["現在庫"]["E2"] = 5
        book["現在庫"]["F2"] = 3
        out = io.BytesIO()
        book.save(out)
        doc = SimpleUploadedFile("数量変更.xlsx", out.getvalue())
        response = self.client.post(
            "/api/import/apply",
            data={"file": doc, "system_name": "制御A-別サイト"},
        )
        self.assertEqual(response.status_code, 201, response.content)
        system = SystemMaster.objects.get(name="制御A-別サイト")
        from .services import balance
        self.assertEqual(balance(system.pk, self.part.pk), 5)
        self.assertEqual(
            StockAlertSetting.objects.get(system=system, part=self.part).threshold, 3
        )

    def test_excel_import_rejects_modified_parts_and_negative_quantity(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        self.movement("IN", 3)
        original = b"".join(
            self.client.get(f"/api/export?system={self.system.pk}").streaming_content
        )
        for column, value in [("E2", -2), ("B2", "FAKE-100")]:
            workbook = load_workbook(io.BytesIO(original))
            workbook["現在庫"][column] = value
            out = io.BytesIO()
            workbook.save(out)
            response = self.client.post(
                "/api/import/apply",
                data={
                    "file": SimpleUploadedFile("改ざん.xlsx", out.getvalue()),
                    "system_name": "取込テスト",
                },
            )
            self.assertEqual(response.status_code, 400, response.content)
            self.assertFalse(SystemMaster.objects.filter(name="取込テスト").exists())

    def test_excel_import_refuses_old_display_only_format(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook
        book = Workbook()
        ws = book.active
        ws.title = "現在庫"
        ws.append(["システム", "品番", "型式", "品名", "現在庫", "下限値", "警告"])
        data = io.BytesIO()
        book.save(data)
        resp = self.client.post(
            "/api/import/preview",
            data={"file": SimpleUploadedFile("旧形式.xlsx", data.getvalue())},
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("メタデータ", resp.json()["error"])

    def test_excel_includes_zero_when_alerted(self):
        replace_alerts(self.system, [{"part_id": self.part.id, "threshold": 1}])
        response = self.client.get("/api/export")
        from openpyxl import load_workbook
        content = b"".join(response.streaming_content)
        ws = load_workbook(io.BytesIO(content)).active
        self.assertEqual(ws["E2"].value, 0)
        self.assertEqual(ws["G2"].value, "下限以下")

    def test_master_api_assigns_optional_equipment_and_category(self):
        response = self.client.post(
            "/api/equipment",
            data=json.dumps({"name": "仕分設備", "alias": "ソータ"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)
        equipment_id = response.json()["item"]["id"]
        section_response = self.client.post(
            "/api/equipment/sections",
            data=json.dumps({"equipment_id": equipment_id, "name": "搬送部"}),
            content_type="application/json",
        )
        self.assertEqual(section_response.status_code, 201)
        section_id = section_response.json()["item"]["id"]
        category = self.client.post(
            "/api/categories",
            data=json.dumps({"name": "機械部品"}),
            content_type="application/json",
        ).json()["item"]["id"]
        result = self.client.post(
            "/api/parts",
            data=json.dumps({
                "part_number": "P200", "model_number": "V1", "name": "ローラー",
                "equipment_id": equipment_id, "section_id": section_id,
                "category_id": category,
            }),
            content_type="application/json",
        )
        self.assertEqual(result.status_code, 201)
        self.assertEqual(result.json()["item"]["section_id"], section_id)

        invalid = self.client.post(
            "/api/parts",
            data=json.dumps({
                "part_number": "P201", "model_number": "V1", "name": "ローラー",
                "section_id": section_id,
            }),
            content_type="application/json",
        )
        self.assertEqual(invalid.status_code, 400)

    def test_api_csrf_and_validation(self):
        client = self.client_class(enforce_csrf_checks=True)
        denied = client.post("/api/history", data="{}", content_type="application/json")
        self.assertEqual(denied.status_code, 403)
        client.get("/api/health")
        token = client.cookies["csrftoken"].value
        allowed = client.post(
            "/api/history", data=json.dumps({
                "system_id": self.system.id, "part_id": self.part.id,
                "direction": "OUT", "quantity": 1, **self.actor
            }), content_type="application/json", HTTP_X_CSRFTOKEN=token,
        )
        self.assertEqual(allowed.status_code, 400)
