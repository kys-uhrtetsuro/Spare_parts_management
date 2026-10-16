"""Portable, system-scoped Excel snapshot export and safe import.

The visible sheet is editable. Master metadata is stored in hidden worksheets.
An import creates opening stock movements; it intentionally does not copy
historical movements from the source database.
"""
import hashlib
import io
import math
import zipfile
from decimal import Decimal
from collections import defaultdict
from datetime import datetime

import pandas as pd
from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook

from .models import (
    EquipmentMaster, EquipmentSection, PartCategory, PartMaster, PartVersion,
    StockAlertSetting, StockHistory, StockImportBatch, SystemMaster, SystemPartPlacement,
)
from .services import InputError, WRITE_LOCK, balances, text_field, reconcile_alert_lineage

FORMAT = "SPARE_PARTS_SNAPSHOT_V4"
SUPPORTED_FORMATS = {"SPARE_PARTS_SNAPSHOT_V2", "SPARE_PARTS_SNAPSHOT_V3", FORMAT}
HEADERS = ("システム", "品番", "型式", "品名", "現在庫", "下限値", "警告")
SCHEMAS = {
    "__manifest": ("key", "value"),
    "__systems": ("id", "name"),
    "__equipment": ("id", "system_id", "name", "alias"),
    "__sections": ("id", "equipment_id", "name"),
    "__categories": ("id", "name"),
    "__parts": ("id", "part_number", "model_number", "name",
                "previous_part_id", "category_id"),
    "__placements": ("system_id", "part_id", "equipment_id", "section_id"),
    # The following definitions are legacy only (V2/V3).
    # V4 does not export or use them, but older workbooks may contain them.
    "__stocks": ("system_id", "part_id", "quantity"),
    "__alerts": ("system_id", "part_id", "threshold"),
}
# Human-readable headers for newly exported metadata. Internal dictionary
# keys remain stable; older workbooks with physical-name headers still import.
EXCEL_HEADERS = {
    "__manifest": ("項目", "値"),
    "__systems": ("システムID", "システム名"),
    "__equipment": ("機器ID", "システムID", "機器名", "機器通称"),
    "__sections": ("部位ID", "機器ID", "部位名"),
    "__categories": ("分類ID", "分類名"),
    "__parts": ("部品ID", "品番", "型式", "品名", "前版部品ID", "分類ID"),
    "__placements": ("システムID", "部品ID", "機器ID", "部位ID"),
    "__stocks": ("システムID", "部品ID", "在庫数"),
    "__alerts": ("システムID", "部品ID", "下限値"),
}
SHEET_LABELS = {
    "__manifest": "ファイル情報",
    "__systems": "システム情報",
    "__equipment": "機器情報",
    "__sections": "部位情報",
    "__categories": "分類情報",
    "__parts": "部品情報",
    "__placements": "部品割当情報",
    "__stocks": "在庫情報",
    "__alerts": "下限設定情報",
}
FIELD_LABELS = {
    ("システム", "name"): "システム名",
    ("機器", "name"): "機器名",
    ("機器", "alias"): "機器通称",
    ("分類", "name"): "分類名",
    ("部位", "name"): "部位名",
    ("部品", "part_number"): "品番",
    ("部品", "model_number"): "型式",
    ("部品", "name"): "品名",
}
MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_INFLATED_BYTES = 40 * 1024 * 1024
MAX_ROWS = 10000


def safe_excel(value):
    """Prevent spreadsheet formula execution in the public-facing sheet."""
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value


def nonnegative(value, label, *, allow_zero=True):
    if isinstance(value, bool) or value is None:
        raise InputError(f"{label}には整数を入力してください")
    if isinstance(value, float):
        if not value.is_integer():
            raise InputError(f"{label}には整数を入力してください")
        value = int(value)
    if isinstance(value, str):
        if not value.strip().isdigit():
            raise InputError(f"{label}には整数を入力してください")
        value = int(value.strip())
    if not isinstance(value, int) or value < (0 if allow_zero else 1) or value > 2147483647:
        raise InputError(f"{label}には{'0' if allow_zero else '1'}以上の整数を入力してください")
    return value


def _select_ids(system_ids, filters):
    """Part masters actively referenced by selected systems, plus their lineage."""
    used = set(StockHistory.objects.filter(system_id__in=system_ids).values_list("part_id", flat=True))
    used.update(StockAlertSetting.objects.filter(system_id__in=system_ids).values_list("part_id", flat=True))
    used.update(SystemPartPlacement.objects.filter(system_id__in=system_ids).values_list("part_id", flat=True))
    if filters:
        categories = PartMaster.objects.all()
        if filters.get("category"):
            categories = categories.filter(category_id=filters["category"])
        allowed = set(categories.values_list("pk", flat=True))
        equipment = {}
        if filters.get("equipment"):
            equipment["equipment_id"] = filters["equipment"]
        if filters.get("section"):
            equipment["section_id"] = filters["section"]
        if equipment:
            allowed.intersection_update(SystemPartPlacement.objects.filter(
                system_id__in=system_ids, **equipment
            ).values_list("part_id", flat=True))
        used.intersection_update(allowed)
    all_parts = dict(PartMaster.objects.values_list("id", "previous_part_id"))
    # Include ancestors and successors so the copied system knows its upgrade chains.
    selected = set(used)
    for _ in range(len(all_parts) + 1):
        expanded = selected.copy()
        for pid in selected:
            parent = all_parts.get(pid)
            if parent:
                expanded.add(parent)
        for pid, parent in all_parts.items():
            if parent in selected:
                expanded.add(pid)
        if expanded == selected:
            break
        selected = expanded
    return selected


def write_snapshot(system_id=None, filters=None):
    filters = filters or {}
    systems = list(SystemMaster.objects.filter(pk=system_id) if system_id else SystemMaster.objects.all())
    system_ids = [obj.pk for obj in systems]
    part_ids = _select_ids(system_ids, filters)
    # Referenced predecessor/successor masters are exported to preserve lineage,
    # but should not secretly expand a classification-filtered stock snapshot.
    parts = list(PartMaster.objects.filter(pk__in=part_ids).order_by("id"))
    # Machine instances are system-owned, even if no part has been attached.
    equipment = list(EquipmentMaster.objects.filter(system_id__in=system_ids).order_by("id"))
    equipment_ids = {e.id for e in equipment}
    sections = list(EquipmentSection.objects.filter(
        equipment_id__in=equipment_ids
    ).order_by("id"))
    placements = list(SystemPartPlacement.objects.filter(
        system_id__in=system_ids, part_id__in=part_ids
    ).order_by("system_id", "part_id"))
    category_ids = {p.category_id for p in parts if p.category_id}
    categories = list(PartCategory.objects.filter(pk__in=category_ids).order_by("id"))
    stock_pairs = {(sid, pid) for sid in system_ids for pid in part_ids}
    if filters:
        category_filter = filters.get("category")
        machine_filter = filters.get("equipment")
        section_filter = filters.get("section")
        matching_placements = {
            (x.system_id, x.part_id) for x in placements
            if (not machine_filter or x.equipment_id == machine_filter)
            and (not section_filter or x.section_id == section_filter)
        }
        stock_pairs = {
            (sid, pid) for sid, pid in stock_pairs
            if (not category_filter or next(p for p in parts if p.pk == pid).category_id == category_filter)
            and (not (machine_filter or section_filter) or (sid, pid) in matching_placements)
        }
    stock = balances()
    alerts = list(StockAlertSetting.objects.filter(
        system_id__in=system_ids, part_id__in=part_ids
    ).order_by("system_id", "part_id"))
    threshold_map = {(a.system_id, a.part_id): a.threshold for a in alerts}
    part_map = {p.id: p for p in parts}
    sys_map = {s.id: s.name for s in systems}
    # Only the visible sheet contains stock quantities and alert thresholds.
    # Hidden sheets contain metadata for restoring equipment and part lineage.
    visible = []
    for sid in sorted(system_ids):
        for pid in sorted(part_ids):
            enabled = (sid, pid) in stock_pairs
            qty = stock.get((sid, pid), 0) if enabled else 0
            threshold = threshold_map.get((sid, pid)) if enabled else None
            if qty > 0 or threshold is not None:
                p = part_map[pid]
                visible.append({
                    "システム": safe_excel(sys_map[sid]),
                    "品番": safe_excel(p.part_number),
                    "型式": safe_excel(p.model_number),
                    "品名": safe_excel(p.name),
                    "現在庫": qty,
                    "下限値": threshold,
                    "警告": "下限以下" if threshold is not None and qty <= threshold else "",
                })

    # Pandas owns the human-facing table; openpyxl owns relational metadata.
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(visible, columns=HEADERS).to_excel(
            writer, sheet_name="現在庫", index=False
        )
        sheets = {
            "__manifest": [("format", FORMAT), ("created_at", datetime.now().isoformat())],
            "__systems": [(s.id, s.name) for s in systems],
            "__equipment": [(x.id, x.system_id, x.name, x.alias) for x in equipment],
            "__sections": [(x.id, x.equipment_id, x.name) for x in sections],
            "__categories": [(x.id, x.name) for x in categories],
            "__parts": [(
                p.id, p.part_number, p.model_number, p.name,
                p.previous_part_id, p.category_id,
            ) for p in parts],
            "__placements": [
                (p.system_id, p.part_id, p.equipment_id, p.section_id) for p in placements
            ],
        }
        for sheetname, rows in sheets.items():
            ws = writer.book.create_sheet(sheetname)
            ws.append(EXCEL_HEADERS[sheetname])
            for row in rows:
                ws.append(row)
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        # Explicit string prevents hidden metadata being interpreted as formulas.
                        cell.data_type = "s"
            ws.sheet_state = "hidden"
        ws = writer.book["現在庫"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, width in {"A": 24, "B": 24, "C": 23, "D": 38, "E": 13, "F": 13, "G": 17}.items():
            ws.column_dimensions[col].width = width
    return output.getvalue()


def cell_value_label(value):
    """Show the original Excel value without exposing internal field names."""
    if value is None or value == "":
        return "空欄"
    kind = ("文字列" if isinstance(value, str)
            else "真偽値" if isinstance(value, bool)
            else "数値" if isinstance(value, (int, float))
            else "その他")
    readable = str(value).replace("\\r", "\\r").replace("\\n", "\\n")
    if len(readable) > 80:
        readable = readable[:80] + "…"
    return f"「{readable}」（{kind}）"


def cell_error(item, field, value, reason):
    sheet = item["__sheet__"]
    column = chr(65 + SCHEMAS[sheet].index(field))
    row = item["__row__"]
    label = EXCEL_HEADERS[sheet][SCHEMAS[sheet].index(field)]
    return InputError(
        f"「{SHEET_LABELS[sheet]}」シート {column}{row}（{label}）、"
        f"入力値：{cell_value_label(value)}。{reason}"
    )


def _table(book, name):
    label = SHEET_LABELS[name]
    if name not in book.sheetnames:
        raise InputError(f"取り込み用メタデータ「{label}」シートがありません")
    ws = book[name]
    if ws.sheet_state == "visible":
        raise InputError(f"取り込み用の「{label}」シートは非表示にしてください")
    rows = ws.iter_rows(values_only=True)
    header = next(rows, ())
    expected = SCHEMAS[name]
    if tuple(header[:len(expected)]) not in (expected, EXCEL_HEADERS[name]):
        raise InputError(f"「{label}」シートの見出しが異なります")
    results = []
    for n, row in enumerate(rows, 2):
        if n > MAX_ROWS:
            raise InputError("Excelの行数が上限を超えています")
        if not any(v is not None for v in row):
            continue
        if len(row) > len(expected) and any(v is not None for v in row[len(expected):]):
            raise InputError(f"「{label}」シートの余分な列を削除してください")
        record = dict(zip(expected, row))
        record["__row__"] = n
        record["__sheet__"] = name
        results.append(record)
    return results


def excel_identifier(value, label, maximum, required=True):
    """Normalize Excel numeric cells before matching source and destination parts.

    Spreadsheet cell formatting alone does not change stored cell types.
    Do not interpret boolean/date cells as identifiers or silently zero-pad IDs.
    """
    if isinstance(value, bool):
        raise InputError(f"{label}には文字列または数値を入力してください")
    if isinstance(value, int):
        value = str(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise InputError(f"{label}には有限の数値を入力してください")
        value = str(int(value)) if value.is_integer() else format(Decimal(str(value)), "f")
    return text_field(value, label, maximum, required)


def _records(rows, required_name, *, columns):
    mapping = {}
    for r in rows:
        raw_id = r.get("id")
        try:
            rid = nonnegative(raw_id, f"{required_name}ID", allow_zero=False)
        except InputError as exc:
            raise cell_error(r, "id", raw_id, str(exc)) from exc
        if rid in mapping:
            raise cell_error(r, "id", raw_id, f"{required_name}IDが重複しています")
        for name, length, mandatory in columns:
            label = FIELD_LABELS[(required_name, name)]
            raw = r.get(name)
            value = raw if raw is not None else ""
            validate = (excel_identifier if required_name == "部品"
                        and name in ("part_number", "model_number") else text_field)
            try:
                r[name] = validate(value, label, length, mandatory)
            except InputError as exc:
                raise cell_error(r, name, raw, str(exc)) from exc
        mapping[rid] = r
    return mapping


def read_snapshot(blob):
    if len(blob) > MAX_XLSX_BYTES:
        raise InputError("ファイルサイズの上限は10MBです")
    if len(blob) < 4 or not blob.startswith(b"PK"):
        raise InputError("xlsx形式のファイルを選択してください")
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            if sum(f.file_size for f in z.infolist()) > MAX_INFLATED_BYTES:
                raise InputError("展開後のファイルサイズが上限を超えています")
        book = load_workbook(io.BytesIO(blob), read_only=False, data_only=True)
    except (zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
        raise InputError("Excelファイルを開けません") from exc
    try:
        manifest = dict((r["key"], r["value"]) for r in _table(book, "__manifest"))
        if manifest.get("format") not in SUPPORTED_FORMATS:
            raise InputError("このExcelは対応する移植形式ではありません")

        systems = _records(_table(book, "__systems"), "システム", columns=[("name", 120, True)])
        equipment = _records(_table(book, "__equipment"), "機器", columns=[
            ("name", 160, True), ("alias", 160, False),
        ])
        categories = _records(_table(book, "__categories"), "分類", columns=[("name", 160, True)])
        sections = _records(_table(book, "__sections"), "部位", columns=[("name", 160, True)])
        parts = _records(_table(book, "__parts"), "部品", columns=[
            ("part_number", 160, True), ("model_number", 160, False),
            ("name", 250, True),
        ])
        if not systems:
            raise InputError("移植対象のシステムがありません")
        if len({r["name"] for r in systems.values()}) != len(systems):
            raise InputError("システム名が重複しています")

        def optref(value, mapping, label, *, item=None, field=None):
            if value is None or value == "":
                return None
            try:
                rid = nonnegative(value, label, allow_zero=False)
            except InputError as exc:
                if item is not None:
                    raise cell_error(item, field, value, str(exc)) from exc
                raise
            if rid not in mapping:
                reason = (
                    f"Excel内の参照先「{label}ID {rid}」が非表示メタデータに存在しません。"
                    "取り込み先DBのIDではなく、Excelファイル内部の参照不整合です。"
                )
                if item is not None:
                    raise cell_error(item, field, value, reason)
                raise InputError(reason)
            return rid

        for machine in equipment.values():
            machine["system_id"] = optref(machine.get("system_id"), systems, "システム", item=machine, field="system_id")
            if machine["system_id"] is None:
                raise InputError("機器にはシステムを指定してください")
        if len({(e["system_id"], e["name"]) for e in equipment.values()}) != len(equipment):
            raise InputError("同一システム内で機器名が重複しています")
        for section in sections.values():
            section["equipment_id"] = optref(section.get("equipment_id"), equipment, "機器", item=section, field="equipment_id")
            if section["equipment_id"] is None:
                raise InputError("部位に機器が指定されていません")
        seen_identifiers = set()
        for part in parts.values():
            pair = (part["part_number"], part["model_number"])
            if pair in seen_identifiers:
                raise InputError("部品の品番と型式が重複しています")
            seen_identifiers.add(pair)
            for col, table, name in [
                ("previous_part_id", parts, "前版部品"), ("category_id", categories, "分類"),
            ]:
                part[col] = optref(part.get(col), table, name, item=part, field=col)

        placements = []
        placement_keys = set()
        for item in _table(book, "__placements"):
            sid = optref(item.get("system_id"), systems, "システム", item=item, field="system_id")
            pid = optref(item.get("part_id"), parts, "部品", item=item, field="part_id")
            eid = optref(item.get("equipment_id"), equipment, "機器", item=item, field="equipment_id")
            sec = optref(item.get("section_id"), sections, "部位", item=item, field="section_id")
            key = (sid, pid, eid, sec)
            if key in placement_keys:
                raise InputError("同一システム・部品・機器・部位の組み合わせが重複しています")
            placement_keys.add(key)
            if sid is None or pid is None or eid is None:
                raise InputError("機器割当の参照が不足しています")
            if equipment[eid]["system_id"] != sid:
                raise InputError("割当機器の所属システムが一致しません")
            if sec is not None and sections[sec]["equipment_id"] != eid:
                raise InputError("割当部位の所属機器が一致しません")
            placements.append((sid, pid, eid, sec))

        # Topological order guards against circular and self-referential versions.
        sorted_parts = []
        pending = set(parts)
        while pending:
            ready = sorted(p for p in pending if parts[p]["previous_part_id"] not in pending)
            if not ready:
                raise InputError("品番・型式の新旧対応が循環しています")
            sorted_parts.extend(ready)
            pending.difference_update(ready)

        # V2/V3 may contain __stocks and __alerts, but they are intentionally
        # ignored: quantities and thresholds must come exclusively from 現在庫.
        # Missing legacy sheets are not an error.
        ws = book["現在庫"] if "現在庫" in book else None
        if ws is None:
            raise InputError("現在庫シートがありません")
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, ())
        if tuple(header[:len(HEADERS)]) != HEADERS:
            raise InputError("現在庫シートの列定義が異なります")
        sys_by_name = {safe_excel(r["name"]): sid for sid, r in systems.items()}
        part_by_identity = {
            (safe_excel(r["part_number"]), safe_excel(r["model_number"])): pid
            for pid, r in parts.items()
        }
        shown = set()
        final_stock = {}
        final_alerts = {}
        next_part_id = min(0, *parts.keys()) - 1
        for rownum, row in enumerate(iterator, 2):
            if rownum > MAX_ROWS:
                raise InputError("現在庫一覧の行数が上限を超えています")
            if not any(v is not None for v in row):
                continue
            vals = dict(zip(HEADERS, row))
            sid = sys_by_name.get(vals.get("システム"))
            # An empty string in the export becomes an empty Excel cell (None).
            model_number = "" if vals.get("型式") is None else vals["型式"]
            try:
                part_number = excel_identifier(
                    vals.get("品番"), "品番", 160
                )
            except InputError as exc:
                raise InputError(
                    f"「現在庫」シート B{rownum}（品番）、"
                    f"入力値：{cell_value_label(vals.get('品番'))}。{exc}"
                ) from exc
            try:
                model_number = excel_identifier(
                    model_number, "型式", 160, False
                )
            except InputError as exc:
                raise InputError(
                    f"「現在庫」シート C{rownum}（型式）、"
                    f"入力値：{cell_value_label(vals.get('型式'))}。{exc}"
                ) from exc
            pid = part_by_identity.get((part_number, model_number))
            if sid is None:
                raise InputError(
                    f"「現在庫」シート A{rownum}（システム）、"
                    f"入力値：{cell_value_label(vals.get('システム'))}。"
                    "Excelのシステム情報と一致しません"
                )
            if pid is None:
                # A new visible row has no hidden ID; allocate an import-local
                # negative ID, never an ID from the destination database.
                try:
                    name = text_field(vals.get("品名"), "品名", 250)
                except InputError as exc:
                    raise InputError(
                        f"「現在庫」シート D{rownum}（品名）、"
                        f"入力値：{cell_value_label(vals.get('品名'))}。{exc}"
                    ) from exc
                pid = next_part_id
                next_part_id -= 1
                parts[pid] = {
                    "part_number": part_number, "model_number": model_number,
                    "name": name, "previous_part_id": None, "category_id": None,
                }
                sorted_parts.append(pid)
                part_by_identity[(part_number, model_number)] = pid
            elif vals.get("品名") != safe_excel(parts[pid]["name"]):
                raise InputError(
                    f"「現在庫」シート D{rownum}（品名）、"
                    f"入力値：{cell_value_label(vals.get('品名'))}。"
                    f"「部品情報」の品名「{parts[pid]['name']}」と一致しません"
                )
            pair = (sid, pid)
            if pair in shown:
                raise InputError(
                    f"「現在庫」シート {rownum}行目（品番・型式）、"
                    f"品番「{part_number}」、型式「{model_number}」が重複しています"
                )
            shown.add(pair)
            try:
                final_stock[pair] = nonnegative(vals.get("現在庫"), "現在庫")
            except InputError as exc:
                raise InputError(
                    f"「現在庫」シート E{rownum}（現在庫）、"
                    f"入力値：{cell_value_label(vals.get('現在庫'))}。{exc}"
                ) from exc
            alert = vals.get("下限値")
            if alert not in (None, ""):
                try:
                    final_alerts[pair] = nonnegative(alert, "下限値", allow_zero=False)
                except InputError as exc:
                    raise InputError(
                        f"「現在庫」シート F{rownum}（下限値）、"
                        f"入力値：{cell_value_label(alert)}。{exc}"
                    ) from exc
            else:
                final_alerts.pop(pair, None)
        return {
            "checksum": hashlib.sha256(blob).hexdigest(),
            "systems": systems, "equipment": equipment, "sections": sections,
            "categories": categories, "parts": parts, "sorted_parts": sorted_parts,
            "stock": final_stock, "alerts": final_alerts, "placements": placements,
        }
    finally:
        book.close()


def available_fork_targets():
    """Any existing system can receive an additive import."""
    return list(SystemMaster.objects.order_by("name", "pk").values("id", "name"))


def preview_snapshot(payload):
    sysnames = [r["name"] for r in payload["systems"].values()]
    existing = set(SystemMaster.objects.filter(name__in=sysnames).values_list("name", flat=True))
    duplicate = StockImportBatch.objects.filter(checksum=payload["checksum"]).exists()
    return {
        "checksum": payload["checksum"],
        "systems": sysnames,
        "system_count": len(sysnames),
        "equipment_count": len(payload["equipment"]),
        "part_count": len(payload["parts"]),
        "opening_rows": sum(v > 0 for v in payload["stock"].values()),
        "quantity_total": sum(payload["stock"].values()),
        "alert_count": len(payload["alerts"]),
        "conflicts": sorted(existing),
        "already_imported": duplicate,
        "available_targets": available_fork_targets() if len(payload["systems"]) == 1 else [],
    }


def apply_snapshot(payload, *, rename=None, target_system_id=None):
    """Atomically fork into new systems or append records to an existing system."""
    if target_system_id is not None:
        from .services import positive_int
        target_system_id = positive_int(target_system_id, "取り込み先システムID")
        if rename is not None:
            raise InputError("既存システムを選ぶ場合、新しいシステム名は指定できません")
        if len(payload["systems"]) != 1:
            raise InputError("既存システムへの取り込みは単一システムのExcelのみ対応します")
    if rename is not None:
        rename = text_field(rename, "移植先システム名", 120)
        if len(payload["systems"]) != 1:
            raise InputError("システム名の変更は単一システムのファイルのみ対応します")
    system_names = {
        sid: rename if rename is not None else r["name"]
        for sid, r in payload["systems"].items()
    }
    with WRITE_LOCK, transaction.atomic():
        if StockImportBatch.objects.filter(checksum=payload["checksum"]).exists():
            raise InputError("このExcelファイルは既に取り込まれています")
        if target_system_id is not None:
            target = SystemMaster.objects.select_for_update().filter(pk=target_system_id).first()
            if target is None:
                raise InputError("取り込み先システムが存在しません")
            source_id = next(iter(payload["systems"]))
            system_map = {source_id: target.pk}
            system_names = {source_id: target.name}
        else:
            if SystemMaster.objects.filter(name__in=system_names.values()).exists():
                raise InputError("同名システムがあります。取り込み先システム名を変更してください")
            system_map = {
                sid: SystemMaster.objects.create(name=name).pk
                for sid, name in system_names.items()
            }
        eq_map = {}
        for src_id, obj in payload["equipment"].items():
            destination_id = system_map[obj["system_id"]]
            # Equipment IDs may differ between exported SQLite databases.
            # Only coalesce records whose system, name AND alias agree.
            equipment, created = EquipmentMaster.objects.get_or_create(
                system_id=destination_id, name=obj["name"],
                defaults={"alias": obj["alias"]},
            )
            if not created and equipment.alias != obj["alias"]:
                raise InputError(
                    f"機器『{obj['name']}』の通称が既存データと異なります。"
                    "異なる機器をIDだけで統合することはできません"
                )
            eq_map[src_id] = equipment.pk

        section_map = {}
        for src_id, obj in payload["sections"].items():
            section, _ = EquipmentSection.objects.get_or_create(
                equipment_id=eq_map[obj["equipment_id"]], name=obj["name"],
            )
            section_map[src_id] = section.pk

        cat_map = {}
        for src_id, obj in payload["categories"].items():
            target, _ = PartCategory.objects.get_or_create(name=obj["name"])
            cat_map[src_id] = target.pk

        part_map = {}
        for src_id in payload["sorted_parts"]:
            obj = payload["parts"][src_id]
            expected = {
                "name": obj["name"],
                "previous_part_id": part_map.get(obj["previous_part_id"]),
                "category_id": cat_map.get(obj["category_id"]),
            }
            target, created = PartMaster.objects.get_or_create(
                part_number=obj["part_number"], model_number=obj["model_number"],
                defaults=expected,
            )
            if not created and any(getattr(target, key) != value for key, value in expected.items()):
                raise InputError(
                    f"部品『{target.part_number} / {target.model_number}』が既存マスタと異なります"
                )
            part_map[src_id] = target.pk
            if obj["previous_part_id"]:
                prior = payload["parts"][obj["previous_part_id"]]
                PartVersion.objects.get_or_create(
                    part_id=part_map[obj["previous_part_id"]],
                    old_part_number=prior["part_number"],
                    old_model_number=prior["model_number"],
                    new_part_number=obj["part_number"],
                    new_model_number=obj["model_number"],
                )

        existing_placements = set(
            SystemPartPlacement.objects.filter(
                system_id__in=system_map.values()
            ).values_list("system_id", "part_id", "equipment_id", "section_id")
        )
        new_placements = []
        for sid, pid, eid, section in payload["placements"]:
            key = (system_map[sid], part_map[pid], eq_map[eid], section_map.get(section))
            if key not in existing_placements:
                new_placements.append(SystemPartPlacement(
                    system_id=key[0], part_id=key[1],
                    equipment_id=key[2], section_id=key[3],
                ))
                existing_placements.add(key)
        SystemPartPlacement.objects.bulk_create(new_placements)
        new_rows = []
        for (sid, pid), qty in payload["stock"].items():
            if qty:
                part_id = part_map[pid]
                part = payload["parts"][pid]
                new_rows.append(StockHistory(
                    system_id=system_map[sid], part_id=part_id,
                    direction=StockHistory.IN, quantity=qty,
                    operator="Excel移行",
                    note=("Excel取り込みによる在庫加算" if target_system_id is not None
                          else "Excel移植による期首在庫"),
                    part_number_snapshot=part["part_number"],
                    model_number_snapshot=part["model_number"],
                ))
        StockHistory.objects.bulk_create(new_rows)
        # Thresholds are settings, not stock quantities: never add them.
        # Preserve existing destination thresholds; add imported settings only
        # for previously unmonitored parts.
        current_alerts = set(StockAlertSetting.objects.filter(
            system_id__in=system_map.values()
        ).values_list("system_id", "part_id"))
        StockAlertSetting.objects.bulk_create([
            StockAlertSetting(
                system_id=system_map[sid], part_id=part_map[pid], threshold=value,
            )
            for (sid, pid), value in payload["alerts"].items()
            if (system_map[sid], part_map[pid]) not in current_alerts
        ])
        # Imported historical alert metadata must follow the same revision rule.
        for new_system_id in system_map.values():
            reconcile_alert_lineage(new_system_id)
        StockImportBatch.objects.create(
            checksum=payload["checksum"], system_count=len(system_map),
            opening_rows=len(new_rows),
        )
        return {"systems": list(system_names.values()), "opening_rows": len(new_rows)}
